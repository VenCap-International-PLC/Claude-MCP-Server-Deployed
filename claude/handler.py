# claude/handler.py
import os
import json
import re
import logging
from datetime import datetime

import anthropic
import chainlit as cl
import plotly.graph_objects as go

from db_config import connect_readonly, connect_insight_readonly
from claude.security import is_safe_query, validate_schema_access
from claude.report_builder import build_report, extract_report_block
from claude.system_prompt import SYSTEM_PROMPT as VENCAP_SYSTEM_PROMPT
from claude.config import CLAUDE_MODEL, MAX_TOKENS

import asyncio

logger = logging.getLogger(__name__)

client = anthropic.AsyncAnthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))


ALLOWED_SCHEMAS = ["bot", "pbi"]

TOOLS = [
    {
        "name": "run_sql",
        "description": (
            "Executes a SELECT query against VenCap SQL Server databases "
            "and returns the results as a list of rows. "
            "Use this for any question about NAV, capital calls, distributions, "
            "fund performance, portfolio exposure, investor transactions, "
            "stock holdings, IPO data, or share valuations. "
            "Only SELECT statements are permitted. "
            "Set database to 'gp' for fund-level queries (bot.* schema) "
            "and stock pipeline queries (pbi.* schema), "
            "or 'lp' for investor-level queries (bot.* schema only)."
        ),
        "input_schema": {
            "type": "object",
            "properties": {
                "query": {
                    "type": "string",
                    "description": "A valid SQL SELECT statement targeting bot.* schema (GP or LP) or pbi.* schema (GP stock pipeline only)."
                },
                "database": {
                    "type": "string",
                    "enum": ["gp", "lp"],
                    "description": "Which database to query: 'gp' for fund-level data (bot.* schema), 'lp' for investor data (bot.* schema)."
                }
            },
            "required": ["query", "database"]
        }
    }
]


# ── SQL execution ─────────────────────────────────────────────────────────────

def run_sql(query: str, database: str = "gp", user_message: str = "") -> str:
    import time

    if not is_safe_query(query):
        logger.warning(f"BLOCKED by keyword check: {query[:200]}")
        return "Error: Query blocked by safety check. Only SELECT statements are permitted."

    if not validate_schema_access(query, ALLOWED_SCHEMAS):
        logger.warning(f"BLOCKED by schema check: {query[:200]}")
        return "Error: Query references a schema you do not have access to."

    conn = None
    try:
        start = time.perf_counter()
        conn = connect_insight_readonly() if database == "lp" else connect_readonly()
        db_label = "LP" if database == "lp" else "GP"

        cursor = conn.cursor()
        cursor.execute(query)
        rows    = cursor.fetchall()
        columns = [col[0] for col in cursor.description] if cursor.description else []
        duration_ms = round((time.perf_counter() - start) * 1000, 1)

        if not rows:
            return "No results found for this query."

        result = [dict(zip(columns, row)) for row in rows]
        _log_query(query, len(rows), db_label, user_message, duration_ms)
        return json.dumps(result, default=str)

    except Exception as e:
        logger.error(f"SQL execution error ({database}): {e}\nQuery: {query[:200]}")
        return f"Database error: {str(e)}"

    finally:
        if conn:
            conn.close()

# ── Audit logger ──────────────────────────────────────────────────────────────

def _log_query(sql: str, row_count: int, database: str = "GP", 
               user_message: str = "", duration_ms: float = None):
    entry = {
        "timestamp":    datetime.now().isoformat(),
        "database":     database,
        "row_count":    row_count,
        "duration_ms":  duration_ms,
        "user_message": user_message,
        "sql":          sql,
    }
    try:
        with open("audit_log.jsonl", "a") as f:
            f.write(json.dumps(entry) + "\n")
    except Exception as e:
        logger.warning(f"Could not write audit log: {e}")
# ── Chart extraction & rendering ──────────────────────────────────────────────

def _extract_all_chart_blocks(text: str) -> tuple[list[dict], str]:
    pattern = r"<chart>(.*?)</chart>"
    matches = re.findall(pattern, text, re.DOTALL)
    charts = []
    for m in matches:
        try:
            charts.append(json.loads(m.strip()))
        except json.JSONDecodeError:
            pass
    cleaned = re.sub(pattern, "", text, flags=re.DOTALL).strip()
    return charts, cleaned


def _build_plotly_figure(chart: dict) -> go.Figure:
    chart_type = chart.get("type", "bar").lower()
    title      = chart.get("title", "")
    x_label    = chart.get("x_label", "")
    y_label    = chart.get("y_label", "")
    x          = chart.get("x", [])
    y          = chart.get("y", [])
    labels     = chart.get("labels", x)
    values     = chart.get("values", y)
    series     = chart.get("series", [])

    fig = go.Figure()

    if chart_type == "pie":
        fig.add_trace(go.Pie(labels=labels, values=values))
    elif series:
        for s in series:
            sx   = s.get("x", x)
            sy   = s.get("y", [])
            name = s.get("name", "")
            if chart_type == "line":
                fig.add_trace(go.Scatter(x=sx, y=sy, mode="lines+markers", name=name))
            elif chart_type == "scatter":
                fig.add_trace(go.Scatter(x=sx, y=sy, mode="markers", name=name))
            else:
                fig.add_trace(go.Bar(x=sx, y=sy, name=name))
    else:
        if chart_type == "line":
            fig.add_trace(go.Scatter(x=x, y=y, mode="lines+markers"))
        elif chart_type == "scatter":
            fig.add_trace(go.Scatter(x=x, y=y, mode="markers"))
        else:
            fig.add_trace(go.Bar(x=x, y=y))

    fig.update_layout(
        title=title,
        xaxis_title=x_label,
        yaxis_title=y_label,
        height=450,
    )
    return fig


# ── Excel export ──────────────────────────────────────────────────────────────

def _extract_excel_block(text: str) -> tuple[list[dict] | None, str]:
    pattern = r"<excel>(.*?)</excel>"
    match = re.search(pattern, text, re.DOTALL)
    if not match:
        return None, text
    try:
        data = json.loads(match.group(1).strip())
    except json.JSONDecodeError:
        return None, text
    cleaned = re.sub(pattern, "", text, flags=re.DOTALL).strip()
    return data, cleaned


def _build_excel(data: list[dict]) -> bytes:
    import io
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VenCap Data"

    if not data:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    headers = list(data[0].keys())
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(header))

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Follow-up suggestions ─────────────────────────────────────────────────────

def _extract_suggestions(text: str) -> tuple[list[str], str]:
    pattern = r"<suggestions>(.*?)</suggestions>"
    match = re.search(pattern, text, re.DOTALL)
    if not match:
        return [], text
    try:
        suggestions = json.loads(match.group(1).strip())
        if not isinstance(suggestions, list):
            return [], text
    except json.JSONDecodeError:
        return [], text
    cleaned = re.sub(pattern, "", text, flags=re.DOTALL).strip()
    return suggestions, cleaned


# ── History sanitiser ─────────────────────────────────────────────────────────

def _sanitise_history(history: list) -> list:
    """
    Remove any tool_result messages that have no matching tool_use
    in the preceding assistant message. Prevents the Anthropic API
    400 error: 'unexpected tool_use_id found in tool_result blocks'.

    Handles three orphan cases:
      1. tool_result is the very first message (clean is empty)
      2. tool_result follows a non-assistant message
      3. tool_result IDs don't match the preceding assistant's tool_use IDs
    """
    clean = []
    for msg in history:
        if msg["role"] == "user" and isinstance(msg["content"], list):
            if any(isinstance(b, dict) and b.get("type") == "tool_result" for b in msg["content"]):
                # No preceding message at all, or preceding message isn't assistant → orphan
                if not clean or clean[-1]["role"] != "assistant":
                    logger.warning("Dropping orphaned tool_result from history (no preceding assistant message).")
                    continue
                prev_content = clean[-1]["content"]
                if not isinstance(prev_content, list):
                    logger.warning("Dropping orphaned tool_result from history (preceding assistant has no content list).")
                    continue
                tool_use_ids = {
                    b.get("id") for b in prev_content
                    if isinstance(b, dict) and b.get("type") == "tool_use"
                }
                result_ids = {
                    b.get("tool_use_id") for b in msg["content"]
                    if isinstance(b, dict) and b.get("type") == "tool_result"
                }
                if not result_ids.issubset(tool_use_ids):
                    logger.warning("Dropping orphaned tool_result from history (IDs do not match preceding tool_use).")
                    continue
        clean.append(msg)
    return clean


# ── Main Claude handler ───────────────────────────────────────────────────────

async def handle_with_claude(user_message: str) -> dict:
    history = cl.user_session.get("claude_history", [])
    history.append({"role": "user", "content": user_message})

    msg = cl.Message(content="⏳ Thinking...")
    await msg.send()

    try:
        # Inject cached query results into system prompt so Claude can reason
        # from prior data without re-querying
        query_cache = cl.user_session.get("query_cache", [])
        if query_cache:
            cache_block = "\n\n━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━\n"
            cache_block += "DATA ALREADY RETRIEVED THIS SESSION — USE THIS, DO NOT RE-QUERY:\n"
            cache_block += "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━\n"
            for entry in query_cache[-5:]:  # last 5 queries only
                cache_block += f"Q: {entry['question']}\nResult: {entry['result'][:2000]}\n\n"
            dynamic_system = VENCAP_SYSTEM_PROMPT + cache_block
        else:
            dynamic_system = VENCAP_SYSTEM_PROMPT

        history = _sanitise_history(history)
        response = await client.messages.create(
            model      = CLAUDE_MODEL,
            max_tokens = MAX_TOKENS,
            system     = dynamic_system,
            tools      = TOOLS,
            messages   = history
        )

        while response.stop_reason == "tool_use":
            tool_results = []
            for block in response.content:
                if block.type == "tool_use" and block.name == "run_sql":
                    generated_sql = block.input.get("query", "")
                    database = block.input.get("database", "gp")
                    logger.info(f"Claude generated SQL ({database.upper()}): {generated_sql[:300]}")

                    msg.content = "🔍 Querying database..."
                    await msg.update()

                    db_result = await asyncio.to_thread(run_sql, generated_sql, database, user_message)

                    # Cache raw results in session so follow-ups don't re-query
                    query_cache = cl.user_session.get("query_cache", [])
                    query_cache.append({
                        "question": user_message,
                        "database": database,
                        "sql": generated_sql,
                        "result": db_result
                    })
                    # Keep last 10 query results only
                    cl.user_session.set("query_cache", query_cache[-10:])
                    tool_results.append({
                        "type":        "tool_result",
                        "tool_use_id": block.id,
                        "content":     db_result
                    })

            history.append({
                "role": "assistant",
                "content": [
                    block.model_dump() if hasattr(block, "model_dump") else block
                    for block in response.content
                ]
            })
            history.append({"role": "user",      "content": tool_results})

            msg.content = "✍️ Formatting answer..."
            await msg.update()
            history = _sanitise_history(history)
            response = await client.messages.create(
                model      = CLAUDE_MODEL,
                max_tokens = MAX_TOKENS,
                system     = dynamic_system,
                tools      = TOOLS,
                messages   = history
            )

        final_answer = ""
        for block in response.content:
            if hasattr(block, "text"):
                final_answer += block.text

        if not final_answer:
            final_answer = "I was unable to generate an answer. Please try rephrasing."

        history.append({"role": "assistant", "content": final_answer})

        # Trim to MAX_HISTORY_TURNS (each turn = 1 user + 1 assistant message)
        from claude.config import MAX_HISTORY_TURNS
        max_messages = MAX_HISTORY_TURNS * 2
        if len(history) > max_messages:
            history = history[-max_messages:]
            # Trimming can cut an assistant tool_use block and leave its
            # tool_result orphaned at the start of the trimmed history.
            # Sanitise immediately so the corrupted state is never saved.
            history = _sanitise_history(history)

        cl.user_session.set("claude_history", history)

        # Remove thinking message
        await msg.remove()

        # Extract all special blocks
        charts_data, text_after_charts = _extract_all_chart_blocks(final_answer)
        excel_data,  text_after_excel  = _extract_excel_block(text_after_charts)
        report_data, text_after_report = extract_report_block(text_after_excel)
        suggestions, clean_text        = _extract_suggestions(text_after_report)

        figures = [_build_plotly_figure(c) for c in charts_data]

        return {
            "text":        clean_text,
            "charts":      figures,
            "excel_data":  excel_data,
            "report_data": report_data,
            "suggestions": suggestions,
        }

    except anthropic.APIError as e:
        await msg.remove()
        logger.error(f"Anthropic API error: {e}")
        return {"text": f"API error: {str(e)}. Please try again.", "charts": [], "excel_data": None, "report_data": None, "suggestions": []}
    except Exception as e:
        await msg.remove()
        logger.error(f"Unexpected error in Claude handler: {e}")
        return {"text": "An unexpected error occurred. Please try again.", "charts": [], "excel_data": None, "report_data": None, "suggestions": []}


# ── Analyse this action callback ──────────────────────────────────────────────

@cl.action_callback("analyse_chart")
async def on_analyse_chart(action: cl.Action):
    chart_json = (action.payload or {}).get("chart_data", "")
    if not chart_json:
        await cl.Message(content="No chart data to analyse.").send()
        return
    prompt = f"Please analyse this data and provide key insights, trends, and anything notable:\n\n{chart_json}"
    result = await handle_with_claude(prompt)
    elements = [cl.Plotly(figure=fig, display="inline", size="large") for fig in result["charts"]]
    await cl.Message(content=result["text"], elements=elements).send()


@cl.action_callback("followup_question")
async def on_followup_question(action: cl.Action):
    question = (action.payload or {}).get("question", "")
    if not question:
        return
    await handle(question)


# ── Public entry point ────────────────────────────────────────────────────────

async def handle(user_message: str, message=None) -> None:
    result = await handle_with_claude(user_message)

    text        = result["text"]
    charts      = result["charts"]
    excel_data  = result["excel_data"]
    report_data = result.get("report_data")
    suggestions = result["suggestions"]

    elements = []
    chart_payloads = []
    for i, fig in enumerate(charts):
        elements.append(cl.Plotly(figure=fig, display="inline", size="large"))
        chart_payloads.append(json.dumps(fig.to_dict(), default=str)[:3000])

    if excel_data:
        try:
            excel_bytes = _build_excel(excel_data)
            elements.append(
                cl.File(
                    content=excel_bytes,
                    name="vencap_export.xlsx",
                    display="inline",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            )
        except Exception as e:
            logger.warning(f"Excel export failed: {e}")
            text += f"\n\n_(Excel export failed: {e})_"

    if report_data:
        try:
            wants_both = report_data.get("type", "word") == "both"
            types = ["word", "pdf"] if wants_both else [report_data.get("type", "word")]
            for rtype in types:
                rd = {**report_data, "type": rtype}
                rbytes, rname, rmime = build_report(rd)
                elements.append(
                    cl.File(
                        content=rbytes,
                        name=rname,
                        display="inline",
                        mime=rmime
                    )
                )
        except Exception as e:
            logger.warning(f"Report generation failed: {e}")
            text += f"\n\n_(Report generation failed: {e})_"

    actions = []
    for i, payload in enumerate(chart_payloads):
        actions.append(
            cl.Action(
                name="analyse_chart",
                label=f"🔍 Analyse chart {i+1}" if len(charts) > 1 else "🔍 Analyse this",
                payload={"chart_data": payload},
            )
        )

    for suggestion in suggestions[:3]:
        actions.append(
            cl.Action(
                name="followup_question",
                label=f"💡 {suggestion}",
                payload={"question": suggestion},
            )
        )

    await cl.Message(content=text, elements=elements, actions=actions).send()


def clear_history() -> None:
    try:
        cl.user_session.set("claude_history", [])
        cl.user_session.set("query_cache", [])
    except Exception:
        pass